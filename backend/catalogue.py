import io
import re
from pathlib import Path
from typing import Dict, List, Optional
import openpyxl
from openpyxl.cell.cell import MergedCell

CAT_DIR = Path(__file__).parent / "catalogues_src"
SUBJECT_ALIASES = {
    "maths": ["mathematics", "maths", "wiskunde"],
    "maths-lit": ["mathematical literacy", "maths lit", "maths literacy", "wiskundige geletterdheid"],
    "physical-sciences": ["physical sciences", "physical science", "fisiese wetenskappe", "science"],
    "natural-sciences": ["natural sciences", "natural science", "natuurwetenskappe", "science"],
    "technical-maths": ["technical mathematics", "technical maths", "tegniese wiskunde"],
    "technical-science": ["technical sciences", "technical science", "tegniese wetenskappe"],
}
SUBJECT_FILES = {
    "maths": ["Maths"], "maths-lit": ["Maths Lit"], "physical-sciences": ["Physical Sciences", "Science"],
    "natural-sciences": ["Natural Sciences"],
    "technical-maths": ["Technical Maths"], "technical-science": ["Technical Science"],
}
FILES = {}
for _g in range(12, 6, -1):
    for _key, _names in SUBJECT_FILES.items():
        _f = next((f"Grade {_g} {n} Catalogue.xlsx" for n in _names if (CAT_DIR / f"Grade {_g} {n} Catalogue.xlsx").exists()), None)
        if _f:
            FILES[(str(_g), _key)] = _f
CATALOGUES = {f"gr{g}-{subj}": (f, g) for (g, subj), f in FILES.items()}
SHEETS = {"National": "DBE", "IEB": "IEB", "KZN": "KZN", "EC": "Eastern Cape", "WC": "Western Cape", "GP": "Gauteng",
          "NW": "North West", "FS": "Free State", "MP": "Mpumalanga", "LP": "Limpopo"}
TITLE_SESSIONS = [("NOV", "Nov"), ("SEPT", "Sep"), ("MAY", "May/Jun"), ("JUN", "May/Jun"), ("FEB", "Feb/Mar"), ("MAR", "Feb/Mar"), ("EXEMPLAR", "Exemplar")]
GREY = "FFD9D9D9"  # rows for exams that were never written


def is_grey(cell) -> bool:
    return cell.fill is not None and cell.fill.fill_type == "solid" and str(cell.fill.start_color.rgb) == GREY


def catalogue_key(meta) -> Optional[str]:
    subj = (meta.subject or "").strip().lower()
    for key, aliases in SUBJECT_ALIASES.items():
        if subj in aliases and (meta.grade, key) in FILES:
            return f"gr{meta.grade}-{key}"
    return None


def _title_session(title: str) -> str:
    up = title.upper().split("-", 1)[-1]
    return next((v for k, v in TITLE_SESSIONS if k in up), "")


def parse_blocks(ws) -> List[dict]:
    anchors = {}
    for rng in ws.merged_cells.ranges:
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                anchors[(row, col)] = (rng.min_row, rng.min_col)

    def val(r, c):
        r0, c0 = anchors.get((r, c), (r, c))
        v = ws.cell(r0, c0).value
        return str(v).strip() if v is not None and str(v).strip() else None

    titles = [r for r in range(1, ws.max_row + 1) if "PAST PAPERS" in str(ws.cell(r, 2).value or "").upper()]
    blocks = []
    for i, t in enumerate(titles):
        title = str(ws.cell(t, 2).value).strip()
        m = re.search(r"\(?P(\d)\)?", title.split("-", 1)[-1])
        cols, errata_col = [], None
        for c in range(3, ws.max_column + 1):
            top = val(t + 1, c)
            sub = val(t + 2, c) if anchors.get((t + 2, c)) != anchors.get((t + 1, c)) or (t + 2, c) not in anchors else None
            if not top:
                continue
            if top.lower() == "errata":
                errata_col = c
                continue
            cols.append((c, f"{top} - {sub}" if sub and sub != top else top))
        end = titles[i + 1] if i + 1 < len(titles) else ws.max_row + 1
        rows, grey = {}, set()
        for r in range(t + 3, end):
            ym = re.search(r"(\d{4})", str(ws.cell(r, 2).value or ""))
            if ym:
                rows[ym.group(1)] = r
                if is_grey(ws.cell(r, 2)):
                    grey.add(ym.group(1))
        blocks.append({"title": title, "session": _title_session(title), "paper": m.group(1) if m else "",
                       "cols": cols, "errata_col": errata_col, "rows": rows, "grey": grey})
    return blocks


_cache: Dict[str, dict] = {}


def structure(key: str) -> dict:
    if key not in _cache:
        wb = openpyxl.load_workbook(CAT_DIR / CATALOGUES[key][0])
        _cache[key] = {ws.title: parse_blocks(ws) for ws in wb.worksheets}
    return _cache[key]


def _find_block(blocks, meta, need_session=True):
    paper = meta.paper or "1"  # single-paper exams fill the P1 block
    for b in blocks:
        if b["paper"] == paper and (not need_session or b["session"] == meta.session) and b["cols"]:
            return b
    return None


def locate(meta) -> dict:
    key = catalogue_key(meta)
    if not key:
        return {"found": False, "reason": "No catalogue for this subject/grade"}
    struct = structure(key)
    sheet = SHEETS.get(meta.province, "Misc")
    if sheet not in struct:
        sheet = "Misc"
    block = _find_block(struct.get(sheet, []), meta)
    base = {"key": key, "workbook": CATALOGUES[key][0], "sheet": sheet}
    topic_block = block or _find_block(struct.get(sheet, []), meta, False) or _find_block(struct.get("DBE", []), meta) \
        or _find_block(struct.get("DBE", []), meta, False)
    topics = [n for _, n in topic_block["cols"]] if topic_block else []
    if not block:
        return {**base, "found": False, "topics": topics, "reason": f"No {meta.session} P{meta.paper} block on sheet {sheet}"}
    row = block["rows"].get(meta.year)
    if not row:
        return {**base, "found": False, "topics": topics, "block": block["title"], "reason": f"No {meta.year} row in {block['title']}"}
    if meta.year in block.get("grey", ()):
        return {**base, "found": False, "topics": topics, "block": block["title"],
                "reason": f"{meta.year} is greyed out in {block['title']} (no such exam on record). Check the paper details"}
    return {**base, "found": True, "topics": topics, "block": block["title"], "row": row, "row_label": f"{meta.session} {meta.year}"}


def topic_cells(questions: List[dict]) -> Dict[str, str]:
    """questions: [{label, topic}] in order. Whole main question -> 'Qn' when all its parts share one topic."""
    mains: Dict[str, List[dict]] = {}
    for q in questions:
        mains.setdefault(q["label"].split(".")[0], []).append(q)
    out: Dict[str, List[str]] = {}
    for main, qs in mains.items():
        topics = {q.get("topic") for q in qs}
        if len(topics) == 1 and None not in topics:
            out.setdefault(qs[0]["topic"], []).append(f"Q{main}")
            continue
        for q in qs:
            if q.get("topic"):
                out.setdefault(q["topic"], []).append(f"Q{q['label']}")
    return {t: ", ".join(v) for t, v in out.items()}


def _write(ws, r, c, value):
    cell = ws.cell(r, c)
    if not isinstance(cell, MergedCell):
        cell.value = value


def build_workbook(key: str, entries: List[dict]) -> bytes:
    """entries: [{meta, questions:[{label, topic}], errata:[labels]}]; later entries win on the same row."""
    wb = openpyxl.load_workbook(CAT_DIR / CATALOGUES[key][0])
    for e in entries:
        loc = locate(e["meta"])
        if not loc.get("found"):
            continue
        ws = wb[loc["sheet"]]
        block = next(b for b in structure(key)[loc["sheet"]] if b["title"] == loc["block"])
        cells = topic_cells(e["questions"])
        for c, name in block["cols"]:
            _write(ws, loc["row"], c, cells.get(name))
        if block["errata_col"]:
            _write(ws, loc["row"], block["errata_col"], ", ".join(f"Q{x}" for x in e["errata"]) or None)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
