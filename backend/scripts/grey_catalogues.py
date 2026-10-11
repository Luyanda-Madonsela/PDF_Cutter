"""One-off: grey out catalogue rows/blocks for exams that were never written. Usage: python scripts/grey_catalogues.py [--apply]"""
import sys
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from catalogue import CAT_DIR, CATALOGUES, parse_blocks, GREY  # noqa: E402

FILL = PatternFill("solid", start_color=GREY, end_color=GREY)
CURRENT_YEAR = 2026
WRITTEN_THIS_YEAR = {"Feb/Mar", "May/Jun"}


def why_not(key, grade, sheet, b, year):
    sess = b["session"]
    if year > CURRENT_YEAR or (year == CURRENT_YEAR and sess not in WRITTEN_THIS_YEAR):
        return "not written yet"
    if year == 2020 and sess == "May/Jun" and not (sheet == "DBE" and grade == 12):
        return "COVID: May/Jun 2020 cancelled"
    if sheet == "IEB":
        return None if grade == 12 and sess == "Nov" else "IEB only sets Gr12 Nov"
    if sheet == "DBE":
        if grade == 12:
            if sess == "Sep":
                return "Gr12 Sept prelims are provincial"
            if sess == "Feb/Mar" and year >= 2018:
                return "Feb/Mar replaced by May/Jun from 2018"
            return None
        return None if sess == "Nov" else "DBE only sets Gr7-11 Nov common papers"
    return None


def has_data(ws, b, r):
    cols = [c for c, _ in b["cols"]] + ([b["errata_col"]] if b["errata_col"] else [])
    return any(ws.cell(r, c).value not in (None, "") for c in cols)


def grey_row(ws, b, r):
    last = max(c for c, _ in b["cols"])
    for c in range(2, last + 1):
        cell = ws.cell(r, c)
        if type(cell).__name__ == "MergedCell":
            continue
        cell.fill = FILL
        f = copy(cell.font)
        cell.font = Font(name=f.name, size=f.size, bold=f.bold, italic=True, color="FF7F7F7F")


def main(apply):
    for key, (fname, grade) in sorted(CATALOGUES.items()):
        wb = openpyxl.load_workbook(CAT_DIR / fname)
        greyed, kept = 0, []
        for ws in wb.worksheets:
            if ws.title in ("Template", "Misc"):
                continue
            for b in parse_blocks(ws):
                for y, r in b["rows"].items():
                    reason = why_not(key, int(grade), ws.title, b, int(y))
                    if not reason:
                        continue
                    if has_data(ws, b, r):
                        kept.append(f"{ws.title} {b['title']} {y} ({reason})")
                        continue
                    grey_row(ws, b, r)
                    greyed += 1
        print(f"{key}: {greyed} rows greyed; {len(kept)} kept because they hold data")
        for k in kept[:5]:
            print("   KEPT", k)
        if apply:
            wb.save(CAT_DIR / fname)


if __name__ == "__main__":
    main("--apply" in sys.argv)
