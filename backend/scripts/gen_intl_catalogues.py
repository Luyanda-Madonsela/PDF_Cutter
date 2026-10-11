"""Generate Cambridge / IB / NBT catalogue workbooks in the CAPS layout. Usage: python scripts/gen_intl_catalogues.py"""
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from catalogue import CAT_DIR, GREY  # noqa: E402
from intl import QUALS, YEARS, file_name  # noqa: E402

HEAD = PatternFill("solid", start_color="FF1F3864", end_color="FF1F3864")
SUBH = PatternFill("solid", start_color="FF2F5597", end_color="FF2F5597")
GREYF = PatternFill("solid", start_color=GREY, end_color=GREY)
THIN = Border(*(Side(style="thin", color="FFBFBFBF"),) * 4)
WHITE = Font(bold=True, color="FFFFFFFF")


def never_written(q, sess, v, year):
    if year > 2026 or (year == 2026 and sess not in ("Feb/Mar", "May/Jun", "Apr")):
        return True
    return q["curriculum"] == "cambridge" and sess == "Feb/Mar" and v != 2  # March series only has variant 2


def block(ws, r, q, sess, sname, p, v):
    topics = q["topics"][p]
    vlabel = f" TZ{v}" if q["curriculum"] == "ib" else f" V{v}" if len(q["variants"]) > 1 else ""
    ws.cell(r, 2, f"{q['label'].upper()} PAST PAPERS - {sname.upper()} (P{p}{vlabel})").font = Font(bold=True, size=13)
    last = 3 + len(topics)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last)
    h1, h2 = r + 1, r + 2
    for c in (2, 3):
        ws.cell(h1, c, "Year" if c == 2 else "Errata")
        ws.merge_cells(start_row=h1, start_column=c, end_row=h2, end_column=c)
    i = 0
    while i < len(topics):
        c = 4 + i
        group, _, sub = topics[i].partition(" - ")
        if sub:
            j = i
            while j < len(topics) and topics[j].startswith(group + " - "):
                ws.cell(h2, 4 + j, topics[j].partition(" - ")[2])
                j += 1
            ws.cell(h1, c, group)
            if j - i > 1:
                ws.merge_cells(start_row=h1, start_column=c, end_row=h1, end_column=3 + j)
            i = j
        else:
            ws.cell(h1, c, topics[i])
            ws.merge_cells(start_row=h1, start_column=c, end_row=h2, end_column=c)
            i += 1
    for rr in (h1, h2):
        for c in range(2, last + 1):
            cell = ws.cell(rr, c)
            if type(cell).__name__ != "MergedCell":
                cell.font, cell.fill, cell.border = WHITE, HEAD if rr == h1 else SUBH, THIN
                cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    rr = h2 + 1
    for y in YEARS:
        grey = never_written(q, sess, v, y)
        for c in range(2, last + 1):
            cell = ws.cell(rr, c, str(y) if c == 2 else None)
            cell.border = THIN
            if grey:
                cell.fill, cell.font = GREYF, Font(italic=True, color="FF7F7F7F")
        rr += 1
    return rr + 2, last


def build(q):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sess, sname in q["series"].items():
        ws = wb.create_sheet(sname)
        ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 10, 12
        r, widest = 2, 3
        for p in q["papers"]:
            for v in q["variants"]:
                r, last = block(ws, r, q, sess, sname, p, v)
                widest = max(widest, last)
        for c in range(4, widest + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 18
        ws.row_dimensions[3].height = 45
    wb.save(CAT_DIR / file_name(q))


if __name__ == "__main__":
    for q in QUALS.values():
        build(q)
        print("built", file_name(q))
