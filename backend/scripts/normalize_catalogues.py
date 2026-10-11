"""One-off: align every catalogue block to the canonical topic columns and normalise question numbers. Usage: python scripts/normalize_catalogues.py [--apply]"""
import re
import sys
import collections
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from catalogue import CAT_DIR, CATALOGUES, parse_blocks  # noqa: E402
from syllabus import CANON, rename  # noqa: E402

TOKEN = re.compile(r"\d+(?:\.\d+)*(?:\([a-z]{1,4}\))*(?:\s*-\s*\d+(?:\.\d+)*(?:\([a-z]{1,4}\))*)?")


def norm_cell(v):
    if v is None or not str(v).strip():
        return None
    s = str(v).strip()
    if re.fullmatch(r"Q?\s*-+", s):
        return "-"
    toks = [re.sub(r"\s+", "", t) for t in TOKEN.findall(s)]
    return ", ".join(f"Q{t}" for t in toks) if toks else s


def canon_lists(key, wb):
    if key in CANON:
        return CANON[key]
    votes = {"1": collections.Counter(), "2": collections.Counter()}
    for ws in wb.worksheets:
        for b in parse_blocks(ws):
            votes.setdefault(b["paper"] or "1", collections.Counter())[tuple(n for _, n in b["cols"])] += 1
    return {p: list(c.most_common(1)[0][0]) for p, c in votes.items() if c}


def rewrite_block(ws, b, new, key, report):
    t = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or "").strip() == b["title"])
    h1, h2 = t + 1, t + 2
    first = min(c for c, _ in b["cols"])
    old_last = max(c for c, _ in b["cols"])
    hdr_style, sub_style = ws.cell(h1, first), ws.cell(h2, first + 0)
    rows = sorted(b["rows"].values())
    data_style = ws.cell(rows[0], first) if rows else hdr_style
    width = ws.column_dimensions[get_column_letter(first)].width
    data = {r: collections.defaultdict(list) for r in rows}
    for c, name in b["cols"]:
        target = rename(key, name)
        for r in rows:
            v = ws.cell(r, c).value
            if v is None or not str(v).strip():
                continue
            if target in new:
                data[r][target].append(str(v))
            else:
                report["lost"].append(f"{ws.title}/{b['title']}/{name}: {v}")
    new_last = first + len(new) - 1
    span = max(old_last, new_last)
    for m in list(ws.merged_cells.ranges):
        if t <= m.min_row <= h2 and (m.min_col >= first or m.min_row == t):
            ws.unmerge_cells(str(m))
    styles = {}
    for c in range(first, span + 1):
        for r in [h1, h2] + rows:
            ws.cell(r, c).value = None
    i = 0
    while i < len(new):
        c = first + i
        group, _, sub = new[i].partition(" - ")
        if sub:
            j = i
            while j < len(new) and new[j].partition(" - ")[0] == group and " - " in new[j]:
                j += 1
            ws.cell(h1, c).value = group
            if j - i > 1:
                ws.merge_cells(start_row=h1, start_column=c, end_row=h1, end_column=first + j - 1)
            for k in range(i, j):
                ws.cell(h2, first + k).value = new[k].partition(" - ")[2]
            i = j
        else:
            ws.cell(h1, c).value = new[i]
            ws.merge_cells(start_row=h1, start_column=c, end_row=h2, end_column=c)
            i += 1
    for k, name in enumerate(new):
        c = first + k
        for r, src in [(h1, hdr_style), (h2, sub_style)] + [(r, data_style) for r in rows]:
            cell = ws.cell(r, c)
            if type(cell).__name__ == "MergedCell":
                continue
            styles[(r, c)] = True
            cell.font, cell.fill, cell.border, cell.alignment = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment)
        if ws.column_dimensions[get_column_letter(c)].width in (None, 0) and width:
            ws.column_dimensions[get_column_letter(c)].width = width
        for r in rows:
            vals = data[r].get(name)
            if vals:
                ws.cell(r, c).value = norm_cell(", ".join(vals))
    if b["errata_col"]:
        for r in rows:
            ws.cell(r, b["errata_col"]).value = norm_cell(ws.cell(r, b["errata_col"]).value)
    ws.merge_cells(start_row=t, start_column=2, end_row=t, end_column=max(new_last, 3))
    return b["cols"] != list(enumerate(new))


def main(apply: bool):
    for key, (fname, _) in sorted(CATALOGUES.items()):
        wb = openpyxl.load_workbook(CAT_DIR / fname)
        lists = canon_lists(key, wb)
        report = {"lost": [], "changed": 0}
        for ws in wb.worksheets:
            for b in parse_blocks(ws):
                new = lists.get(b["paper"] or "1")
                if not new or not b["cols"]:
                    continue
                if [n for _, n in b["cols"]] != new:
                    report["changed"] += 1
                rewrite_block(ws, b, new, key, report)
        print(f"{key}: {report['changed']} blocks re-headed, {len(report['lost'])} values dropped")
        for line in report["lost"][:10]:
            print("   LOST", line)
        if apply:
            wb.save(CAT_DIR / fname)


if __name__ == "__main__":
    main("--apply" in sys.argv)
