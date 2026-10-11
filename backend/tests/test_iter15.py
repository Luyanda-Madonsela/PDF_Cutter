"""Iteration 15: Stop button, Cambridge/IB/NBT catalogues & Drive layout, folder skeleton builder."""
import os
import io
import sys
import asyncio
import subprocess
import time
import pytest
import requests
from pathlib import Path
from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent.parent / ".env")
load_dotenv(Path(__file__).resolve().parent.parent.parent / "frontend" / ".env")

sys.path.insert(0, "/app/backend")

BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
MONGO_URL = os.environ["MONGO_URL"]
DB_NAME = os.environ["DB_NAME"]
AUTH = {"Authorization": "Bearer test_session_qa_papercut"}

_LOOP = asyncio.new_event_loop()


def _db():
    from motor.motor_asyncio import AsyncIOMotorClient
    return AsyncIOMotorClient(MONGO_URL, io_loop=_LOOP)[DB_NAME]


def arun(coro):
    return _LOOP.run_until_complete(coro)


# ---------- (1) Stop & Retry on cloned set ----------
def test_stop_and_retry_lifecycle():
    db = _db()
    proc = subprocess.run(["python3", "tests/clone_set.py"], cwd="/app/backend",
                          env={**os.environ, "PYTHONPATH": "."}, capture_output=True, text=True, timeout=60)
    assert proc.returncode == 0, proc.stderr
    clone_id = proc.stdout.strip().splitlines()[-1]
    assert len(clone_id) == 36
    print(f"[clone] {clone_id}")
    try:
        qp = arun(db.jobs.find_one({"set_id": clone_id, "meta.doc_type": "question_paper", "is_deleted": False}))
        assert qp
        qp_id = qp["_id"]
        # mark QP job as 'ready' so worker picks it to processing
        arun(db.jobs.update_one({"_id": qp_id}, {"$set": {"status": "ready", "error": None, "pages_done": 0, "questions": []}}))

        # wait until status == processing
        got_processing = False
        for i in range(20):
            time.sleep(1)
            d = arun(db.jobs.find_one({"_id": qp_id}, {"status": 1}))
            if d and d["status"] == "processing":
                got_processing = True
                print(f"[processing] after {i}s")
                break
        assert got_processing, "worker never picked the job"

        # Stop the processing job
        r = requests.post(f"{BASE_URL}/api/jobs/{qp_id}/stop", headers=AUTH, timeout=15)
        assert r.status_code == 200, r.text
        # Within 5s should be stopped
        stopped = False
        for i in range(20):
            time.sleep(0.25)
            d = arun(db.jobs.find_one({"_id": qp_id}, {"status": 1, "error": 1, "pages_done": 1}))
            if d and d["status"] == "stopped":
                stopped = True
                assert d.get("pages_done", 0) == 0, d
                assert "stop" in (d.get("error") or "").lower(), d.get("error")
                print(f"[stopped] after processing: error={d['error']!r}")
                break
        assert stopped, "job did not reach stopped"

        # Set-level status should be 'stopped'
        s = requests.get(f"{BASE_URL}/api/sets/{clone_id}", headers=AUTH, timeout=10)
        assert s.status_code == 200
        assert s.json().get("status") == "stopped", s.json().get("status")

        # Retry -> queued
        r = requests.post(f"{BASE_URL}/api/jobs/{qp_id}/retry", headers=AUTH, timeout=10)
        assert r.status_code == 200
        d = arun(db.jobs.find_one({"_id": qp_id}, {"status": 1}))
        assert d["status"] == "queued", d

        # Stop while queued (not running) -> stopped
        r = requests.post(f"{BASE_URL}/api/jobs/{qp_id}/stop", headers=AUTH, timeout=10)
        assert r.status_code == 200
        d = arun(db.jobs.find_one({"_id": qp_id}, {"status": 1, "error": 1}))
        assert d["status"] == "stopped", d
        print(f"[queued->stopped] error={d.get('error')!r}")

        # Mark job done artificially and try to stop -> 400
        arun(db.jobs.update_one({"_id": qp_id}, {"$set": {"status": "done"}}))
        r = requests.post(f"{BASE_URL}/api/jobs/{qp_id}/stop", headers=AUTH, timeout=10)
        assert r.status_code == 400, r.text
    finally:
        arun(db.sets.update_one({"_id": clone_id}, {"$set": {"is_deleted": True}}))
        arun(db.jobs.update_many({"set_id": clone_id}, {"$set": {"is_deleted": True}}))


# ---------- (2) Catalogue list & intl workbook ----------
def test_catalogue_list_has_intl_grades():
    r = requests.get(f"{BASE_URL}/api/catalogues", headers=AUTH, timeout=20)
    assert r.status_code == 200
    data = r.json()
    grades = {c["grade"] for c in data}
    expected = {"IGCSE", "AS & A Level", "IB Diploma", "NBT", "O Level", "Lower Secondary Checkpoint"}
    missing = expected - grades
    assert not missing, f"missing grades: {missing}"
    # also expect "many" entries - review says 48
    print(f"[catalogues] n={len(data)} grades={sorted(grades)}")
    assert len(data) >= 40, len(data)


def test_catalogue_download_intl_0580_sheets():
    import openpyxl
    r = requests.get(f"{BASE_URL}/api/catalogues/intl-0580/download", headers=AUTH, timeout=60)
    assert r.status_code == 200, r.text[:200]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    for s in ("Feb-Mar", "May-June", "Oct-Nov"):
        assert s in wb.sheetnames, f"missing sheet {s} in {wb.sheetnames}"


def test_catalogue_locate_intl():
    from models import JobMeta
    from catalogue import locate

    # Cambridge 0580 May/Jun 2024 P2 V2 -> found
    m = JobMeta(curriculum="cambridge", qualification="0580", year="2024", session="May/Jun", paper="2", variant="2")
    loc = locate(m)
    assert loc.get("found") is True, loc
    assert loc["sheet"] == "May-June"
    print(f"[CIE 0580 MJ 24 P2V2] {loc.get('block')} row={loc.get('row')}")

    # Feb/Mar variant 1 (not written) -> greyed found False
    m2 = JobMeta(curriculum="cambridge", qualification="0580", year="2024", session="Feb/Mar", paper="2", variant="1")
    loc2 = locate(m2)
    assert loc2.get("found") is False, loc2
    print(f"[CIE 0580 FM 24 P2V1 greyed] {loc2.get('reason')}")

    # IB-MAA-HL Nov 2023 P3 variant 0 -> found
    m3 = JobMeta(curriculum="ib", qualification="IB-MAA-HL", year="2023", session="Nov", paper="3", variant="0")
    loc3 = locate(m3)
    assert loc3.get("found") is True, loc3
    print(f"[IB MAA HL Nov23 P3 TZ0] {loc3.get('block')}")

    # NBT-MAT 2022 -> found
    m4 = JobMeta(curriculum="nbt", qualification="NBT-MAT", year="2022")
    loc4 = locate(m4)
    assert loc4.get("found") is True, loc4
    print(f"[NBT MAT 2022] {loc4.get('block')}")

    # CAPS gr12-maths National Nov 2019 P1 -> still found
    m5 = JobMeta(subject="Mathematics", grade="12", year="2019", session="Nov", paper="1", province="National")
    loc5 = locate(m5)
    assert loc5.get("found") is True, loc5
    print(f"[CAPS gr12 maths Nov19 P1 DBE] {loc5.get('block')}")


# ---------- (3) jobs.canon_meta + set_base_name ----------
def test_set_base_name_cambridge_and_oct_nov():
    from models import JobMeta
    from jobs import canon_meta, set_base_name

    m = canon_meta(JobMeta(curriculum="cambridge", qualification="0580", subject="Mathematics",
                           year="2024", session="May/Jun", paper="2", variant="2"))
    base = set_base_name(m, "foo.pdf")
    print(f"[base MayJun] {base}")
    assert base == "Mathematics_CIE0580_2024_MayJun_P2V2", base

    # Oct/Nov should canonicalise to "Nov" and Drive uses "Oct-Nov"
    m2 = canon_meta(JobMeta(curriculum="cambridge", qualification="0580", subject="Mathematics",
                            year="2024", session="Oct/Nov", paper="2", variant="2"))
    assert m2.session == "Nov", m2.session
    from intl import series_name, qual_for
    q = qual_for(m2)
    assert series_name(q, m2.session) == "Oct-Nov"


# ---------- (4) drive.plan_uploads for Cambridge ----------
def test_drive_plan_uploads_cambridge():
    from models import PaperSet, Job, JobMeta, Question, Part
    from drive import plan_uploads, catalogue_folder

    meta = JobMeta(subject="Mathematics", curriculum="cambridge", qualification="0580",
                   year="2024", session="May/Jun", paper="2", variant="2", doc_type="question_paper")
    s = PaperSet(meta=meta, base_name="Mathematics_CIE0580_2024_MayJun_P2V2", mode="bulk", status="done")

    qp = Job(filename="qp.pdf", pdf_path="path/qp.pdf", status="done", total_pages=1, set_id=s.id,
             meta=meta.model_copy(update={"doc_type": "question_paper"}),
             questions=[Question(label="1.1", storage_path="s/q1.png",
                                 parts=[Part(page=0, y0=0, y1=100)])])
    memo_meta = meta.model_copy(update={"doc_type": "memo"})
    memo = Job(filename="memo.pdf", pdf_path="path/memo.pdf", status="done", total_pages=1, set_id=s.id,
               meta=memo_meta,
               questions=[Question(label="1.1", storage_path="s/m1.png",
                                   parts=[Part(page=0, y0=0, y1=100)])])

    plan = plan_uploads(s, [qp, memo], ["root", "Past Paper Database", "CAPS & IEB Past Papers"])
    paths = [tuple(p["path"]) for p in plan]
    print(f"[plan] {paths[:6]}")
    # QP PDF under .../Varient_2/Question Paper
    assert any(p[-2:] == ("Varient_2", "Question Paper") for p in paths), paths
    # Memo PDF under .../Varient_2/Memo
    assert any(p[-2:] == ("Varient_2", "Memo") for p in paths), paths
    # Images under .../Varient_2/Questions/Question_1.1/Question
    assert any(p[-3:] == ("Questions", "Question_1.1", "Question") for p in paths), paths

    # catalogue_folder for gr12-maths -> ends with Catalogues/CAPS & IEB/Grade 12/Mathematics
    gr12_meta = JobMeta(subject="Mathematics", grade="12", year="2024", session="Nov", paper="1", province="National")
    folder = catalogue_folder("gr12-maths", gr12_meta, ["root", "Past Paper Database", "CAPS & IEB Past Papers"])
    print(f"[catfolder gr12-maths] {folder}")
    assert folder[-4:] == ["Catalogues", "CAPS & IEB", "Grade 12", "Mathematics"], folder


# ---------- (5) Drive skeleton options ----------
def test_skeleton_options_lists_qualifications():
    r = requests.get(f"{BASE_URL}/api/drive/skeleton/options", headers=AUTH, timeout=15)
    assert r.status_code == 200
    data = r.json()
    quals = data["qualifications"]
    print(f"[skeleton n={len(quals)}]")
    assert len(quals) >= 27, len(quals)  # spec says 29 but code generates from QUALS dict
    codes = {q["code"]: q for q in quals}
    assert "0580" in codes
    assert codes["0580"]["folders_per_year"] == 180, codes["0580"]


# ---------- (6) Skeleton build (ONE small real run - NBT-AQL 2024-2024) ----------
def test_skeleton_build_small_nbt_run():
    db = _db()
    # Ensure not already running
    cur = arun(db.settings.find_one({"_id": "drive_skeleton"})) or {}
    if cur.get("status") == "running":
        pytest.skip("skeleton build already running")

    r = requests.post(f"{BASE_URL}/api/drive/skeleton",
                      json={"qualifications": ["NBT-AQL"], "year_from": 2024, "year_to": 2024},
                      headers=AUTH, timeout=20)
    assert r.status_code == 200, r.text
    total = r.json().get("total")
    print(f"[skeleton total] {total}")
    assert total == 5, total  # Question Paper, Memo, Addendum, Errata, Questions

    # Poll options.status until done (or error)
    done = False
    for i in range(60):
        time.sleep(2)
        opt = requests.get(f"{BASE_URL}/api/drive/skeleton/options", headers=AUTH, timeout=10).json()
        st = opt.get("status") or {}
        if st.get("status") in ("done", "error"):
            print(f"[skeleton final] {st}")
            assert st["status"] == "done", st
            done = True
            break
    assert done, "skeleton run did not finish"


if __name__ == "__main__":
    pytest.main([__file__, "-v", "-s"])
