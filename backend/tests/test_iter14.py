"""Iteration 14: AI monthly spending limit, school-made tests, Drive sub-question folders."""
import os
import asyncio
import subprocess
import time
import pytest
import requests
from pathlib import Path
from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent.parent / ".env")
load_dotenv(Path(__file__).resolve().parent.parent.parent / "frontend" / ".env")

BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
MONGO_URL = os.environ["MONGO_URL"]
DB_NAME = os.environ["DB_NAME"]
AUTH = {"Authorization": "Bearer test_session_qa_papercut"}
IEB_SET_ID = "4c41f0c1-f54f-49ec-b5c6-db27470415b7"

_LOOP = asyncio.new_event_loop()


def _db():
    from motor.motor_asyncio import AsyncIOMotorClient
    return AsyncIOMotorClient(MONGO_URL, io_loop=_LOOP)[DB_NAME]


def arun(coro):
    return _LOOP.run_until_complete(coro)


# ---------------- (1) Budget endpoint ----------------
def _restore_limit_null():
    r = requests.put(f"{BASE_URL}/api/ai/budget", json={"monthly_limit": None}, headers=AUTH, timeout=10)
    assert r.status_code == 200


@pytest.fixture(autouse=True)
def _always_restore_budget():
    yield
    try:
        requests.put(f"{BASE_URL}/api/ai/budget", json={"monthly_limit": None}, headers=AUTH, timeout=10)
    except Exception:
        pass


def test_budget_get_shape():
    r = requests.get(f"{BASE_URL}/api/ai/budget", headers=AUTH, timeout=10)
    assert r.status_code == 200
    d = r.json()
    for k in ("monthly_limit", "spent", "unpriced_models", "paused_jobs", "month", "reached"):
        assert k in d, f"missing key {k} in {d}"
    assert isinstance(d["spent"], (int, float))
    assert isinstance(d["unpriced_models"], list)
    assert isinstance(d["paused_jobs"], int)
    assert d["month"]  # YYYY-MM


def test_budget_put_tiny_limit_reached_and_blocks_ai_test():
    r = requests.put(f"{BASE_URL}/api/ai/budget", json={"monthly_limit": 0.0001}, headers=AUTH, timeout=10)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["monthly_limit"] == 0.0001
    assert body["reached"] is True

    # /api/ai/test should 400 mentioning 'spending limit'
    t = requests.post(f"{BASE_URL}/api/ai/test", headers=AUTH, timeout=30)
    assert t.status_code == 400, f"expected 400, got {t.status_code}: {t.text[:200]}"
    detail = (t.json().get("detail") or "").lower()
    assert "spending limit" in detail or "limit" in detail, f"unexpected: {detail!r}"


def test_budget_put_negative_rejected():
    r = requests.put(f"{BASE_URL}/api/ai/budget", json={"monthly_limit": -1}, headers=AUTH, timeout=10)
    assert r.status_code == 400


def test_budget_null_removes_limit_and_ai_test_works():
    r = requests.put(f"{BASE_URL}/api/ai/budget", json={"monthly_limit": None}, headers=AUTH, timeout=10)
    assert r.status_code == 200
    assert r.json()["monthly_limit"] is None
    t = requests.post(f"{BASE_URL}/api/ai/test", headers=AUTH, timeout=60)
    # AI /test calls the LLM; a success is 200, but if emergent is temporarily failing we just need "not budget"
    assert t.status_code in (200, 400, 500, 502, 503), t.status_code
    if t.status_code == 400:
        detail = (t.json().get("detail") or "").lower()
        assert "spending limit" not in detail, f"should not be budget-blocked: {detail}"


# ---------------- Unpriced model blocking ----------------
def test_unpriced_model_blocked():
    db = _db()
    # Snapshot current ai settings so we can restore
    before = arun(db.settings.find_one({"_id": "ai"})) or {}
    try:
        # Set limit
        assert requests.put(f"{BASE_URL}/api/ai/budget", json={"monthly_limit": 100}, headers=AUTH, timeout=10).status_code == 200
        # Configure unpriced openrouter model + fake key
        r = requests.put(f"{BASE_URL}/api/ai/settings",
                         json={"provider": "openrouter", "model": "some/unpriced-model", "key": "sk-or-fake"},
                         headers=AUTH, timeout=10)
        assert r.status_code == 200, r.text
        t = requests.post(f"{BASE_URL}/api/ai/test", headers=AUTH, timeout=30)
        assert t.status_code == 400, f"got {t.status_code}: {t.text[:200]}"
        detail = (t.json().get("detail") or "").lower()
        assert "no price" in detail or "price" in detail, f"unexpected detail: {detail!r}"
    finally:
        # Clear openrouter key, back to emergent, limit null
        requests.put(f"{BASE_URL}/api/ai/settings", json={"provider": "openrouter", "key": ""}, headers=AUTH, timeout=10)
        requests.put(f"{BASE_URL}/api/ai/settings", json={"provider": "emergent"}, headers=AUTH, timeout=10)
        _restore_limit_null()


# ---------------- (2) Pause / Resume via worker on CLONED set ----------------
def test_pause_resume_on_cloned_set():
    db = _db()
    # Clone the IEB set
    proc = subprocess.run(
        ["python3", "tests/clone_set.py"],
        cwd="/app/backend", env={**os.environ, "PYTHONPATH": "."},
        capture_output=True, text=True, timeout=60,
    )
    assert proc.returncode == 0, f"clone failed: {proc.stderr}"
    clone_id = proc.stdout.strip().splitlines()[-1]
    print(f"[clone] new set id = {clone_id}")
    assert len(clone_id) == 36

    try:
        # Find QP job of the clone and set it to 'ready'
        qp = arun(db.jobs.find_one({"set_id": clone_id, "meta.doc_type": "question_paper", "is_deleted": False}))
        assert qp, "no QP job in clone"
        qp_id = qp["_id"]
        arun(db.jobs.update_one({"_id": qp_id}, {"$set": {"status": "ready", "error": None, "paused_from": None}}))

        # Apply tiny limit; worker should pause within ~60s
        assert requests.put(f"{BASE_URL}/api/ai/budget", json={"monthly_limit": 0.0001}, headers=AUTH, timeout=10).status_code == 200

        paused = False
        for _ in range(90):
            time.sleep(1)
            doc = arun(db.jobs.find_one({"_id": qp_id}, {"status": 1, "paused_from": 1, "error": 1}))
            if doc and doc.get("status") == "paused":
                paused = True
                print(f"[pause] after {_}s: paused_from={doc.get('paused_from')} error={(doc.get('error') or '')[:80]}")
                assert doc.get("paused_from") == "ready"
                assert "limit" in (doc.get("error") or "").lower()
                break
        assert paused, "QP job never moved to paused"

        # Set-level status should also report paused
        s = requests.get(f"{BASE_URL}/api/sets/{clone_id}", headers=AUTH, timeout=10)
        assert s.status_code == 200
        assert s.json().get("status") == "paused", f"set status: {s.json().get('status')}"

        # Lift limit; job should go back to ready/processing/done
        assert requests.put(f"{BASE_URL}/api/ai/budget", json={"monthly_limit": None}, headers=AUTH, timeout=10).status_code == 200
        resumed = False
        for _ in range(30):
            time.sleep(1)
            doc = arun(db.jobs.find_one({"_id": qp_id}, {"status": 1}))
            if doc and doc.get("status") in ("ready", "queued", "processing", "waiting", "classifying", "done"):
                resumed = True
                print(f"[resume] after {_}s: status={doc['status']}")
                break
        assert resumed, "QP job did not resume after lifting the limit"
    finally:
        # Soft-delete the clone and all its jobs
        arun(db.sets.update_one({"_id": clone_id}, {"$set": {"is_deleted": True}}))
        arun(db.jobs.update_many({"set_id": clone_id}, {"$set": {"is_deleted": True}}))
        _restore_limit_null()


# ---------------- (2) School tests: pure python ----------------
def test_catalogue_locate_school():
    import sys
    sys.path.insert(0, "/app/backend")
    from models import JobMeta
    from catalogue import locate, build_workbook
    import openpyxl, io

    m = JobMeta(subject="Mathematics", grade="12", year="2024", paper="1",
                school="Westville", assessment="Test", term="2")
    loc = locate(m)
    assert loc["found"] is True
    assert loc["sheet"] == "School Tests"
    assert loc.get("school") is True
    print(f"[locate school] {loc}")

    # build_workbook contains School Tests sheet with row
    entry = {"meta": m, "questions": [{"label": "1.1", "topic": None}], "errata": []}
    data = build_workbook("gr12-maths", [entry])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "School Tests" in wb.sheetnames
    ws = wb["School Tests"]
    found_row = False
    for row in ws.iter_rows(values_only=True):
        if row and "Westville" in [str(c) for c in row if c is not None]:
            found_row = True
            print(f"[school row] {row}")
            break
    assert found_row


def test_set_base_name_school_and_drive_folder():
    import sys
    sys.path.insert(0, "/app/backend")
    from models import JobMeta
    from jobs import set_base_name
    from drive import paper_folder

    m = JobMeta(subject="Mathematics", grade="12", year="2024", paper="1",
                school="Westville", assessment="Test", term="2")
    base = set_base_name(m, "foo.pdf")
    print(f"[base] {base}")
    assert "Westville" in base
    assert "T2" in base
    assert "Test" in base
    assert base.endswith("P1")

    folder = paper_folder(m, ["root"])
    print(f"[folder] {folder}")
    assert folder[-2:] == ["Misc", "Westville"]

    # Non-school still maps to Session/Province folders
    m2 = JobMeta(subject="Mathematics", grade="12", year="2024", paper="1",
                 session="Nov", province="IEB")
    folder2 = paper_folder(m2, ["root"])
    print(f"[folder non-school] {folder2}")
    assert "Misc" not in folder2 or folder2[-2] != "Misc"
    # Last two should be session folder + province folder
    assert folder2[-1] in ("IEB", "National", "DBE") or len(folder2) >= 5


# ---------------- (3) Drive plan_uploads: Question_{label} sub-folders ----------------
def test_plan_uploads_question_label_folders():
    import sys
    sys.path.insert(0, "/app/backend")
    from models import PaperSet, Job
    from drive import plan_uploads

    db = _db()
    s_doc = arun(db.sets.find_one({"_id": IEB_SET_ID}))
    assert s_doc
    s = PaperSet.from_mongo(s_doc)
    jobs_docs = arun(db.jobs.find({"set_id": IEB_SET_ID, "is_deleted": False}).to_list(100))
    jobs = [Job.from_mongo(j) for j in jobs_docs]

    plan = plan_uploads(s, jobs, ["root"])
    # Collect question folder names
    qfolders = set()
    for item in plan:
        for part in item.get("path", []):
            if part.startswith("Question_"):
                qfolders.add(part)
    print(f"[qfolders] {sorted(qfolders)[:10]}")
    # At least one sub-question folder like Question_1.1
    sub_qs = [f for f in qfolders if "." in f]
    assert sub_qs, f"no sub-question folders: {qfolders}"

    # Verify Hints and Video entries exist for each question folder
    hints_dirs = {tuple(item["path"]) for item in plan if item["path"][-1] == "Hints"}
    video_dirs = {tuple(item["path"]) for item in plan if item["path"][-1] == "Video"}
    assert hints_dirs
    assert video_dirs
    # And each Hints path's parent is a Question_* folder
    for p in list(hints_dirs)[:3]:
        assert p[-2].startswith("Question_"), p


# ---------------- PATCH set meta w/ school fields on cloned set ----------------
def test_patch_set_meta_school_renames():
    db = _db()
    proc = subprocess.run(
        ["python3", "tests/clone_set.py"],
        cwd="/app/backend", env={**os.environ, "PYTHONPATH": "."},
        capture_output=True, text=True, timeout=60,
    )
    assert proc.returncode == 0, proc.stderr
    clone_id = proc.stdout.strip().splitlines()[-1]
    try:
        payload = {
            "subject": "Physical Sciences", "grade": "12", "year": "2023", "paper": "2",
            "school": "Westville", "assessment": "ControlTest", "term": "2",
            "doc_type": "question_paper"
        }
        r = requests.patch(f"{BASE_URL}/api/sets/{clone_id}/meta", json=payload, headers=AUTH, timeout=30)
        assert r.status_code == 200, r.text
        # Verify base_name now has Westville + T2
        s = arun(db.sets.find_one({"_id": clone_id}))
        base = s.get("base_name", "")
        print(f"[patch base] {base}")
        assert "Westville" in base
        assert "T2" in base
        # Catalogues still works
        c = requests.get(f"{BASE_URL}/api/catalogues", headers=AUTH, timeout=30)
        assert c.status_code == 200
    finally:
        arun(db.sets.update_one({"_id": clone_id}, {"$set": {"is_deleted": True}}))
        arun(db.jobs.update_many({"set_id": clone_id}, {"$set": {"is_deleted": True}}))
