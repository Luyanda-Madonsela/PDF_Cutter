"""Iteration 7 tests: catalogue normalisation + grey rows, refresh-topics, manual erase,
Drive OAuth / status / send plan, ai.cut_bands.
"""
import io
import os
import re
import sys
import time
import pytest
import requests
import openpyxl

sys.path.insert(0, "/app/backend")

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/") or "https://papercut-sa.preview.emergentagent.com"
API = f"{BASE_URL}/api"
QA = "test_session_qa_papercut"
H = {"Authorization": f"Bearer {QA}"}

TOPIC_RE = re.compile(r"^(Q\d+(\.\d+)*(\([a-z]+\))*(-\d+(\.\d+)*)?(, )?)+$")


# ---- helpers ----
def _get(path, **kw):
    return requests.get(f"{API}{path}", headers=H, **kw)


def _post(path, **kw):
    return requests.post(f"{API}{path}", headers=H, **kw)


# ============================================================
# 1) Catalogues: list 21, download, structure, topic cell regex
# ============================================================
@pytest.fixture(scope="module")
def catalogues_list():
    r = _get("/catalogues")
    assert r.status_code == 200, r.text
    data = r.json()
    assert len(data) == 21, f"expected 21 catalogues, got {len(data)}"
    return data


def test_catalogues_list_shape(catalogues_list):
    for c in catalogues_list:
        assert c["filename"].endswith(".xlsx")
        assert c["grade"] in {"7", "8", "9", "10", "11", "12"}


def test_every_catalogue_downloads_openpyxl(catalogues_list):
    """All 21 download endpoints return a valid xlsx."""
    for c in catalogues_list:
        r = _get(f"/catalogues/{c['key']}/download")
        assert r.status_code == 200, f"{c['key']}: {r.status_code}"
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames, c["key"]


def test_catalogue_blocks_consistent_topic_cols_and_regex(catalogues_list):
    """Within each workbook, every block with the same paper number has the same topic columns.
    All non-empty topic cells match the Q-ref regex or '-'."""
    from catalogue import structure  # type: ignore
    bad = []
    for c in catalogues_list:
        struct = structure(c["key"])
        # Per paper number -> set of column-name tuples (must be the same across sheets+sessions)
        by_paper = {}
        for sheet, blocks in struct.items():
            for b in blocks:
                if not b["cols"]:
                    continue
                names = tuple(n for _, n in b["cols"])
                by_paper.setdefault(b["paper"], set()).add(names)
        for p, variants in by_paper.items():
            if len(variants) > 1:
                bad.append(f"{c['key']} P{p}: {variants}")
    assert not bad, "Inconsistent topic columns within a workbook:\n" + "\n".join(bad)


def test_topic_cells_match_regex():
    """Download gr12-maths and gr12-physical-sciences; every non-empty topic cell matches the regex or is '-'."""
    bad = []
    for key in ("gr12-maths", "gr12-physical-sciences"):
        r = _get(f"/catalogues/{key}/download")
        assert r.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        from catalogue import parse_blocks  # type: ignore
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for b in parse_blocks(ws):
                for col, _ in b["cols"]:
                    for yr, row in b["rows"].items():
                        v = ws.cell(row, col).value
                        if v is None:
                            continue
                        s = str(v).strip()
                        if s == "" or s == "-":
                            continue
                        if not TOPIC_RE.match(s):
                            bad.append(f"{key}/{sheet_name}/{b['title']}/{yr}/col{col}: {s!r}")
    assert not bad, "Topic cells failing regex:\n" + "\n".join(bad[:30])


# ============================================================
# 2) locate(): grey rows & IEB
# ============================================================
def test_locate_grey_row_2028():
    from catalogue import locate  # type: ignore
    from models import JobMeta  # type: ignore
    meta = JobMeta(subject="Mathematics", grade="12", year="2028", session="Nov", province="National", paper="1")
    r = locate(meta)
    assert r["found"] is False
    assert "greyed out" in (r.get("reason") or "").lower()


def test_locate_ieb_nov_2019_found():
    from catalogue import locate  # type: ignore
    from models import JobMeta  # type: ignore
    meta = JobMeta(subject="Mathematics", grade="12", year="2019", session="Nov", province="IEB", paper="1")
    r = locate(meta)
    assert r["found"] is True, r


# ============================================================
# 3) refresh-topics
# ============================================================
def test_refresh_topics_endpoint():
    r = _post("/catalogues/refresh-topics")
    assert r.status_code == 200, r.text
    body = r.json()
    assert "updated" in body and "cleared" in body
    assert isinstance(body["updated"], int) and isinstance(body["cleared"], int)
    # Every QP question topic in done sets should be null or in its catalogue topics
    sets = _get("/sets").json()["sets"]
    bad = []
    for s in sets:
        if s["status"] != "done":
            continue
        detail = _get(f"/sets/{s['id']}").json()
        topics = (detail.get("catalogue") or {}).get("topics") or []
        if not topics:
            continue
        qp = next((d for d in detail["documents"] if d["meta"]["doc_type"] == "question_paper"), None)
        if not qp:
            continue
        for q in qp["questions"]:
            t = q.get("topic")
            if t and t not in topics:
                bad.append(f"{s['base_name']} Q{q['label']}: '{t}' not in {topics}")
    assert not bad, bad


# ============================================================
# 4) Manual erase API
# ============================================================
@pytest.fixture(scope="module")
def ieb_set():
    sets = _get("/sets").json()["sets"]
    cand = [s for s in sets if s["status"] == "done" and s["base_name"].startswith("PhysicalSciences_Gr12_2023_Nov_IEB_P2")]
    if not cand:
        pytest.skip("IEB PS set not seeded")
    return _get(f"/sets/{cand[0]['id']}").json()


def _qp_doc(detail):
    return next(d for d in detail["documents"] if d["meta"]["doc_type"] == "question_paper")


def test_erase_get_returns_manual_and_pages(ieb_set):
    qp = _qp_doc(ieb_set)
    q = qp["questions"][0]
    r = _get(f"/jobs/{qp['id']}/questions/{q['id']}/erase")
    assert r.status_code == 200, r.text
    body = r.json()
    assert "manual" in body and "pages" in body
    assert isinstance(body["manual"], list)
    assert isinstance(body["pages"], list) and body["pages"]


def test_manual_erase_invalid_length_rejected(ieb_set):
    qp = _qp_doc(ieb_set)
    q = qp["questions"][0]
    r = _post(f"/jobs/{qp['id']}/questions/{q['id']}/erase",
              json={"boxes": [], "manual": [[1, 100, 100, 200]]})  # 4 ints, not 5
    assert r.status_code == 400, r.text


def test_manual_erase_out_of_range_rejected(ieb_set):
    qp = _qp_doc(ieb_set)
    q = qp["questions"][0]
    page = q["parts"][0]["page"]
    r = _post(f"/jobs/{qp['id']}/questions/{q['id']}/erase",
              json={"boxes": [], "manual": [[page, 0, 0, 2000, 500]]})  # y1>1000
    assert r.status_code == 400, r.text


def test_manual_erase_non_overlapping_rejected(ieb_set):
    """Pick a page where the question has no parts (or box far outside y-range) -> must 400."""
    qp = _qp_doc(ieb_set)
    q = qp["questions"][0]
    page = q["parts"][0]["page"]
    y0 = q["parts"][0]["y0"]
    y1 = q["parts"][0]["y1"]
    # Box fully above the crop
    box_y0 = max(0, y0 - 50)
    box_y1 = max(1, y0 - 10)
    if box_y1 - box_y0 < 2:
        pytest.skip("crop too close to top")
    r = _post(f"/jobs/{qp['id']}/questions/{q['id']}/erase",
              json={"boxes": [], "manual": [[page, box_y0, 100, box_y1, 500]]})
    assert r.status_code == 400, r.text


def test_manual_erase_valid_applies_and_stores(ieb_set):
    """Valid manual box overlapping the crop -> erase_status=applied, erase_manual contains it, storage_path changes."""
    sid = ieb_set["id"]
    detail = _get(f"/sets/{sid}").json()
    qp = _qp_doc(detail)
    # Prefer a question whose erase is currently rejected or none, so we can see transition
    q = next((x for x in qp["questions"] if x["erase_status"] != "applied"), qp["questions"][-1])
    orig_storage = q["storage_path"]
    page = q["parts"][0]["page"]
    y0 = q["parts"][0]["y0"]
    y1 = q["parts"][0]["y1"]
    box = [page, y0 + 5, 100, min(1000, y0 + 25), 300]
    r = _post(f"/jobs/{qp['id']}/questions/{q['id']}/erase",
              json={"boxes": [], "manual": [box]})
    assert r.status_code == 200, r.text
    body = r.json()
    updated = next(x for x in next(d for d in body["documents"] if d["id"] == qp["id"])["questions"] if x["id"] == q["id"])
    assert updated["erase_status"] == "applied", updated
    assert box in updated["erase_manual"], updated["erase_manual"]
    assert updated["storage_path"] != orig_storage


def test_erase_empty_boxes_and_manual_rejected_state(ieb_set):
    """POST with boxes=[] and manual=[] -> status 'rejected' and erase_manual empty."""
    sid = ieb_set["id"]
    detail = _get(f"/sets/{sid}").json()
    qp = _qp_doc(detail)
    q = qp["questions"][1]
    r = _post(f"/jobs/{qp['id']}/questions/{q['id']}/erase", json={"boxes": [], "manual": []})
    assert r.status_code == 200, r.text
    body = r.json()
    updated = next(x for x in next(d for d in body["documents"] if d["id"] == qp["id"])["questions"] if x["id"] == q["id"])
    assert updated["erase_status"] == "rejected"
    assert updated["erase_manual"] == []


# ============================================================
# 5) Drive API
# ============================================================
def test_drive_status_not_connected():
    r = _get("/drive/status")
    assert r.status_code == 200
    body = r.json()
    assert body["connected"] is False
    assert "root" in body and body["root"].startswith("My Drive/")


def test_drive_connect_returns_oauth_url():
    r = _get("/drive/connect")
    assert r.status_code == 200
    url = r.json()["authorization_url"]
    assert url.startswith("https://accounts.google.com/")
    assert "access_type=offline" in url
    assert "redirect_uri=" in url
    # redirect_uri must point at our callback
    assert "api%2Fdrive%2Fcallback" in url or "/api/drive/callback" in url


def test_drive_send_400_when_not_connected():
    sets = _get("/sets").json()["sets"]
    done = next((s for s in sets if s["status"] == "done"), None)
    assert done, "need a done set"
    r = _post(f"/drive/sets/{done['id']}/send")
    assert r.status_code == 400, r.text
    assert "Connect Google Drive" in r.text


def test_drive_callback_bogus_state_redirects_error():
    r = requests.get(f"{API}/drive/callback", params={"state": "bogus", "code": "x"},
                     headers=H, allow_redirects=False)
    assert r.status_code in (302, 307), r.status_code
    loc = r.headers.get("location", "")
    assert "drive=error" in loc, loc


def test_drive_endpoints_require_auth():
    """All /api/drive/* except /callback need auth. Use a fresh Session to bypass the
    module-level monkey-patch in test_papercut_api.py that auto-adds auth headers."""
    sess = requests.Session()
    for path in ("/drive/status", "/drive/connect"):
        r = sess.get(f"{API}{path}")
        assert r.status_code == 401, f"{path} -> {r.status_code}"
    r = sess.post(f"{API}/drive/disconnect")
    assert r.status_code == 401


# ============================================================
# 6) drive.plan_uploads() unit test
# ============================================================
def test_plan_uploads_structure():
    import asyncio
    import drive  # type: ignore
    from jobs import load_set, set_jobs  # type: ignore
    sets = _get("/sets").json()["sets"]
    s = next((x for x in sets if x["status"] == "done" and x["base_name"].startswith("PhysicalSciences_Gr12_2023_Nov_IEB_P2")), None)
    if not s:
        pytest.skip("IEB PS set not seeded")

    async def _load():
        return await load_set(s["id"]), await set_jobs(s["id"])
    try:
        loop = asyncio.get_event_loop()
        if loop.is_closed():
            raise RuntimeError
    except RuntimeError:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
    ps, jobs = loop.run_until_complete(_load())
    root_str = _get("/drive/status").json()["root"]
    root = [p for p in root_str.replace("My Drive/", "", 1).split("/") if p]
    items = drive.plan_uploads(ps, jobs, root)
    assert items
    # Expect: PDFs in Paper_2/Memo_Paper_2/Addendum/Errata; questions under Questions/Paper_2/Question_{main}/...
    paths = ["/".join(it["path"]) for it in items]
    joined = "\n".join(paths)
    assert any("Paper_2" in p and p.endswith("Paper_2") and it["name"].endswith(".pdf") for it, p in zip(items, paths))
    assert any("Memo_Paper_2" in p for p in paths)
    # Question_1/Question subfolder
    assert any("Questions/Paper_2/Question_1/Question" in p for p in paths), joined[:1000]
    # Hints and Video empty folders (name=None)
    hints = [it for it in items if it["name"] is None and "Hints" in it["path"]]
    videos = [it for it in items if it["name"] is None and "Video" in it["path"]]
    assert hints and videos
    # Grade & subject & session segment in paths
    assert any("Grade 12" in p for p in paths)
    assert any("Physical Sciences" in p for p in paths)
    assert any("/Nov/IEB/" in p or p.endswith("Nov/IEB") or "/Nov/IEB" in p for p in paths)
    # PDFs: replace=False
    pdfs = [it for it in items if it.get("mime") == "application/pdf"]
    assert pdfs and all(it["replace"] is False for it in pdfs)
    # PNGs: replace=True
    pngs = [it for it in items if it.get("mime") == "image/png"]
    assert pngs and all(it["replace"] is True for it in pngs)


# ============================================================
# 7) ai.cut_bands unit test
# ============================================================
def test_cut_bands_simple():
    from ai import cut_bands  # type: ignore
    seg = {"label": "1.1", "y0": 100, "y1": 400, "marks": 7, "review": None, "topic": "T"}
    pieces = cut_bands([seg], [(200, 250), (380, 450)])
    assert len(pieces) == 2
    assert (pieces[0]["y0"], pieces[0]["y1"]) == (100, 200)
    assert (pieces[1]["y0"], pieces[1]["y1"]) == (250, 380)
    assert all(p["label"] == "1.1" for p in pieces)
    assert pieces[0]["marks"] == 7
    # marks only on first piece
    assert pieces[1]["marks"] is None
    assert pieces[1]["review"] is None
